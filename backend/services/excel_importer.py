"""Excel → SQLite importer for rental management data."""
from pathlib import Path
from typing import Optional
import openpyxl
from sqlmodel import Session, select

from models import Contract

EXCEL_PATH = Path(__file__).parent.parent.parent / "data" / "rental.xlsx"

SKIP_SHEETS = {"통합", "센터별렌탈현황"}

# 새 컬럼 구조: 모델명이 index 4에 추가됨
COL = {
    "no": 1,
    "item": 2,
    "status": 3,
    "model_name": 4,
    "serial_no": 5,   # KT/LGU 시트에서는 전화번호
    "vendor": 6,
    "corporation": 7,
    "contract_months": 8,
    "start_date": 9,
    "end_date": 10,
    "usage_months": 11,
    "address": 12,
    "monthly_fee": 13,
    "payment_method": 16,
    "notes": 17,
}


def _clean_str(v) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def _clean_date(v) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    return s[:10]


def _clean_fee(v) -> Optional[float]:
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    try:
        f = float(s)
        return f if f > 0 else None
    except ValueError:
        return None


def import_excel(session: Session) -> int:
    """Clear existing data and re-import from Excel. Returns row count."""
    existing = session.exec(select(Contract)).all()
    for c in existing:
        session.delete(c)
    session.commit()

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    count = 0

    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue

        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=4, values_only=True):
            if row[COL["no"]] is None or str(row[COL["no"]]).strip() == "":
                continue
            no_val = _clean_str(row[COL["no"]])
            if no_val and not no_val.replace(".", "").isdigit():
                continue

            safe = lambda i: row[i] if i < len(row) else None

            contract = Contract(
                no=no_val,
                item=_clean_str(safe(COL["item"])),
                status=_clean_str(safe(COL["status"])),
                model_name=_clean_str(safe(COL["model_name"])),
                serial_no=_clean_str(safe(COL["serial_no"])),
                vendor=_clean_str(safe(COL["vendor"])) or sheet_name,
                corporation=_clean_str(safe(COL["corporation"])),
                contract_months=_clean_str(safe(COL["contract_months"])),
                start_date=_clean_date(safe(COL["start_date"])),
                end_date=_clean_date(safe(COL["end_date"])),
                usage_months=_clean_str(safe(COL["usage_months"])),
                address=_clean_str(safe(COL["address"])),
                monthly_fee=_clean_fee(safe(COL["monthly_fee"])),
                payment_method=_clean_str(safe(COL["payment_method"])),
                notes=_clean_str(safe(COL["notes"])),
                source_sheet=sheet_name,
            )
            session.add(contract)
            count += 1

    session.commit()
    return count
